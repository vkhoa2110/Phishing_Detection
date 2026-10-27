"""Build the Vietnamese project report from local evidence; never load .env or call LLMs."""
from pathlib import Path
import argparse
import json
import re
import hashlib
import shutil
from collections import OrderedDict
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
REPORT_SOURCE = ROOT / 'docs/reports/2026-10-26'
OUTPUT = ROOT / 'artifacts/report-builds/2026-10-26'
ASSETS = OUTPUT / 'assets'
NAME = 'Bao_cao_Phishing_ML_LLM_RAG_Le_Van_Khoa'
REFS = {
 'bcp': ('Báo Điện tử Chính phủ', 'Dự báo năm 2026, mức độ tinh vi của lừa đảo trực tuyến gia tăng', '07/01/2026', 'https://baochinhphu.vn/du-bao-nam-2026-muc-do-tinh-vi-cua-cac-hinh-thuc-lua-dao-truc-tuyen-se-tang-102260107101756364.htm'),
 'safe': ('Google', 'Safe Browsing APIs: Overview', '', 'https://developers.google.com/safe-browsing'),
 'vt': ('VirusTotal', 'How it works', '', 'https://docs.virustotal.com/docs/how-it-works'),
 'data': ('T. N. Vu', 'PhishVN: A Time-Stamped Vietnamese URL Phishing Dataset with Impersonation-Scenario Labels and Confidence Tiers, Mendeley Data, phiên bản 4, DOI: 10.17632/b97hxbxtpd.4', '2026', 'https://data.mendeley.com/datasets/b97hxbxtpd/4'),
 'local': ('L. V. Khoa', 'PhishLens: mã nguồn và hồ sơ thực nghiệm của đề tài; bản làm việc được đối chiếu ngày 26/10/2026', '2026', 'https://github.com/vkhoa2110/Phishing_Detection'),
 'xgb': ('T. Chen và C. Guestrin', 'XGBoost: A Scalable Tree Boosting System, KDD 2016, DOI: 10.1145/2939672.2939785', '2016', 'https://arxiv.org/abs/1603.02754'),
 'baseline': ('T. N. Vu', 'PhishVN: Code for the Data in Brief article, bản phát hành dib-113195; train_url_baseline.py', '2026', 'https://github.com/vuthainguyen1602/phishvn/tree/dib-113195'),
 'rag': ('P. Lewis và cộng sự', 'Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks, NeurIPS 2020', '2020', 'https://arxiv.org/abs/2005.11401'),
 'e5paper': ('L. Wang và cộng sự', 'Multilingual E5 Text Embeddings: A Technical Report', '2024', 'https://arxiv.org/abs/2402.05672'),
 'cld': ('Chống Lừa Đảo', 'Về chúng tôi', '', 'https://chongluadao.vn/about'),
 'usercentric': ('A. H. Al Barwani, A. A. Korba và R. W. Anwar', 'User-Centric Phishing Detection: A RAG and LLM-Based Approach', '2026', 'https://arxiv.org/abs/2601.21261'),
 'multi': ('Y. Xue, E. Spero, M. W. Woo, W. Gao và G. Russello', 'MultiPhishGuard: An Explainable and Adaptive Multi-Agent LLM System for Phishing Email Detection, phiên bản 2', '2026', 'https://arxiv.org/abs/2505.23803v2'),
 'fptapi': ('FPT Cloud', 'Call the Inference API: Serverless Inference', '', 'https://docs.fptcloud.com/docs/fpt-ai-factory/services/serverless-inference/tutorials/call-the-api/'),
 'datasheet': ('T. N. Vu', 'Datasheet — PhishVN URL Phishing Dataset, bản phát hành dib-113195', '2026', 'https://github.com/vuthainguyen1602/phishvn/blob/dib-113195/docs/datasheet.md'),
 'e5': ('intfloat', 'multilingual-e5-small: Model card', '', 'https://huggingface.co/intfloat/multilingual-e5-small'),
 'faiss': ('Meta FAIR', 'Faiss Documentation', '', 'https://faiss.ai/'),
 'owasp': ('OWASP', 'LLM Prompt Injection Prevention Cheat Sheet', '', 'https://cheatsheetseries.owasp.org/cheatsheets/LLM_Prompt_Injection_Prevention_Cheat_Sheet.html'),
 'pricing': ('FPT AI Factory', 'Pricing: Serverless Inference — DeepSeek-V4-Flash', '', 'https://factory.fpt.ai/pricing'),
}

def field(p, instruction, cached=''):
    r = OxmlElement('w:r'); f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'),'begin');r.append(f);p._p.append(r)
    r = OxmlElement('w:r');t=OxmlElement('w:instrText');t.set(qn('xml:space'),'preserve');t.text=instruction;r.append(t);p._p.append(r)
    r = OxmlElement('w:r');f=OxmlElement('w:fldChar');f.set(qn('w:fldCharType'),'separate');r.append(f);p._p.append(r)
    if cached:p.add_run(cached)
    r=OxmlElement('w:r');f=OxmlElement('w:fldChar');f.set(qn('w:fldCharType'),'end');r.append(f);p._p.append(r)

def hyperlink(p, text, url, size=10):
    rel=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
    h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel)
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr')
    col=OxmlElement('w:color');col.set(qn('w:val'),'1D4D71');pr.append(col)
    sz=OxmlElement('w:sz');sz.set(qn('w:val'),str(size*2));pr.append(sz)
    r.append(pr);t=OxmlElement('w:t');t.text=text;r.append(t);h.append(r);p._p.append(h)

class Report:
    def __init__(self):
        self.doc=Document();self.refs=OrderedDict();self.captions=[];self.bookmarks=0
        self.setup()
    def setup(self):
        sec=self.doc.sections[0];self.geometry(sec)
        st=self.doc.styles['Normal'];st.font.name='Times New Roman';st.font.size=Pt(13)
        st.element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
        st.paragraph_format.line_spacing=1.3;st.paragraph_format.space_after=Pt(6)
        st.paragraph_format.first_line_indent=Cm(.75)
        for name,size in [('Heading 1',16),('Heading 2',14),('Heading 3',13)]:
            s=self.doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.bold=True
            s.font.color.rgb=RGBColor.from_string('172A3A');s.paragraph_format.first_line_indent=Cm(0)
            s.paragraph_format.space_before=Pt(12);s.paragraph_format.space_after=Pt(8)
        self.doc.styles['Heading 1'].paragraph_format.page_break_before=True
        s=self.doc.styles.add_style('Front Title',WD_STYLE_TYPE.PARAGRAPH)
        s.base_style=self.doc.styles['Heading 1']
        pr=s.element.get_or_add_pPr();ol=OxmlElement('w:outlineLvl');ol.set(qn('w:val'),'9');pr.append(ol)
        for name in ['Caption','TOC 1','TOC 2']:
            if name not in self.doc.styles:
                self.doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH).base_style=self.doc.styles['Normal']
            sty=self.doc.styles[name]
            sty.font.name='Times New Roman';sty.font.size=Pt(12 if name=='Caption' else 11)
            sty.paragraph_format.first_line_indent=Cm(0)
            if name!='Caption':
                sty.paragraph_format.line_spacing=1.05;sty.paragraph_format.space_after=Pt(3)
                sty.paragraph_format.space_before=Pt(0)
            else:
                sty.font.color.rgb=RGBColor.from_string('263746');sty.font.bold=False
        self.doc.core_properties.title='Phát hiện và phân tích nguy cơ lừa đảo trực tuyến bằng ML, LLM và RAG'
        self.doc.core_properties.author='Lê Văn Khoa'
        self.doc.core_properties.subject='Báo cáo dự án công nghệ; hiện trạng và thực nghiệm đến 26/10/2026'
        self.doc.core_properties.keywords='phishing, PhishVN, XGBoost, LLM, RAG, tiếng Việt'
        settings=self.doc.settings.element;u=OxmlElement('w:updateFields');u.set(qn('w:val'),'true');settings.append(u)
    def geometry(self,s):
        s.page_width=Cm(21);s.page_height=Cm(29.7)
        s.top_margin=Cm(2);s.bottom_margin=Cm(2);s.left_margin=Cm(3);s.right_margin=Cm(2)
        s.header_distance=Cm(.9);s.footer_distance=Cm(1)
    def section(self,roman=False):
        s=self.doc.add_section(WD_SECTION_START.NEW_PAGE);self.geometry(s)
        s.header.is_linked_to_previous=False;s.footer.is_linked_to_previous=False
        p=s.header.paragraphs[0];p.paragraph_format.first_line_indent=Cm(0)
        r=p.add_run('LÊ VĂN KHOA · PHÁT HIỆN VÀ PHÂN TÍCH PHISHING');r.font.size=Pt(9);r.font.color.rgb=RGBColor(100,100,100)
        p=s.footer.paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
        field(p,' PAGE ')
        n=OxmlElement('w:pgNumType');n.set(qn('w:start'),'1');n.set(qn('w:fmt'),'lowerRoman' if roman else 'decimal');s._sectPr.append(n)
    def resolve(self,text):
        def sub(m):
            key=m.group(1)
            if key not in REFS:raise KeyError(key)
            if key not in self.refs:self.refs[key]=len(self.refs)+1
            return f'[{self.refs[key]}]'
        return re.sub(r'\[\[([a-z0-9_]+)\]\]',sub,text)
    def p(self,text='',bold=False,small=False):
        p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
        r=p.add_run(self.resolve(text));r.bold=bold
        if small:r.font.size=Pt(11)
        return p
    def h(self,text,level=2):
        if getattr(self,'frontmode',False) and level==1:
            return self.doc.add_paragraph(text,'Front Title')
        return self.doc.add_heading(text,level=level)
    def bullet(self,text):
        p=self.p('• '+text);p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.left_indent=Cm(.35)
    def bookmark(self,p):
        self.bookmarks+=1;name=f'caption_{self.bookmarks}'
        a=OxmlElement('w:bookmarkStart');a.set(qn('w:id'),str(self.bookmarks));a.set(qn('w:name'),name)
        z=OxmlElement('w:bookmarkEnd');z.set(qn('w:id'),str(self.bookmarks));p._p.insert(0,a);p._p.append(z)
        return name
    def caption(self,text,kind):
        p=self.doc.add_paragraph(self.resolve(text),'Caption');p.paragraph_format.first_line_indent=Cm(0)
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER if kind=='Hình' else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.keep_with_next=kind=='Bảng'
        self.captions.append((kind,text,self.bookmark(p)))
    def table(self,caption,headers,rows,widths=None,size=10.5):
        self.caption(caption,'Bảng')
        t=self.doc.add_table(rows=1,cols=len(headers));t.autofit=False
        if widths is None:widths=[16/len(headers)]*len(headers)
        for c,w in zip(t.columns,widths):c.width=Cm(w)
        for c,h in zip(t.rows[0].cells,headers):c.text=self.resolve(str(h))
        for row in rows:
            for c,v in zip(t.add_row().cells,row):c.text=self.resolve(str(v))
        for i,row in enumerate(t.rows):
            trpr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trpr.append(no)
            if i==0:
                rep=OxmlElement('w:tblHeader');trpr.append(rep)
            for j,c in enumerate(row.cells):
                c.width=Cm(widths[j]);pr=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
                for tag,val in [('top','70'),('bottom','70'),('left','85'),('right','85')]:
                    x=OxmlElement('w:'+tag);x.set(qn('w:w'),val);x.set(qn('w:type'),'dxa');mar.append(x)
                pr.append(mar)
                if i==0 or i%2==0:
                    sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E5EDF3' if i==0 else 'F4F6F8');pr.append(sh)
                for p in c.paragraphs:
                    p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1.05
                    p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(2)
                    if len(rows)<=7 and sum(len(str(v)) for row_ in rows for v in row_)<2200:
                        p.paragraph_format.keep_with_next=i<len(t.rows)-1
                    for r in p.runs:r.font.size=Pt(size);r.bold=i==0
        p=self.doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.space_before=Pt(0)
        p.paragraph_format.line_spacing=1;p.add_run().font.size=Pt(1)
    def fig(self,name,caption,width=16):
        p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
        p.paragraph_format.keep_with_next=True;p.add_run().add_picture(str(ASSETS/name),width=Cm(width))
        self.caption(caption,'Hình')
    def code(self,text):
        p=self.doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.line_spacing=1.05
        p.paragraph_format.space_after=Pt(8)
        r=p.add_run(text);r.font.name='Liberation Mono';r.font.size=Pt(9)
        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F1F3F5');p._p.get_or_add_pPr().append(sh)
    def front(self):
        self.frontmode=True
        for text,size,bold in [('ĐẠI HỌC QUỐC GIA HÀ NỘI',14,True),('TRƯỜNG ĐẠI HỌC CÔNG NGHỆ',14,True),('KHOA CÔNG NGHỆ THÔNG TIN',13,True),('──────────',13,False),('',13,False),('BÁO CÁO DỰ ÁN CÔNG NGHỆ',18,True),('',13,False),('PHÁT HIỆN VÀ PHÂN TÍCH\nNGUY CƠ LỪA ĐẢO TRỰC TUYẾN\nBẰNG HỌC MÁY, LLM VÀ RAG',19,True),('(Phishing Detection and Analysis\nusing Machine Learning, LLMs and RAG)',13,False)]:
            p=self.doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.first_line_indent=Cm(0)
            p.paragraph_format.space_after=Pt(10);r=p.add_run(text);r.font.size=Pt(size);r.bold=bold
        self.p('')
        for text in ['Sinh viên thực hiện: Lê Văn Khoa','Mã sinh viên: 23020092','Lớp: K68I-IT2','Giảng viên hướng dẫn: ThS. Trần Ngọc Trúc Linh']:
            p=self.p(text);p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.left_indent=Cm(2)
        p=self.p('Hà Nội, tháng 10 năm 2026');p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_before=Pt(24)
        self.section(roman=True)
        self.h('THÔNG TIN VỀ BẢN BÁO CÁO',1).paragraph_format.page_break_before=False
        self.p('Báo cáo ghi nhận hiện trạng dự án và các kết quả có thể kiểm chứng đến ngày 26/10/2026. Cấu trúc được tham khảo từ mẫu học phần Dự án công nghệ và đề cương của sinh viên; nội dung được cập nhật theo mã nguồn, tệp cấu hình và kết quả thực nghiệm đang lưu trong dự án.')
        self.p('Các mục được ghi “đã triển khai” hoặc “đã đo” có bằng chứng cài đặt hay kết quả tương ứng. Các mục được ghi “thiết kế”, “mục tiêu” hoặc “kế hoạch” chưa được coi là tính năng hay kết quả đã hoàn thành. Riêng RAG đã có kho tài liệu, chưa có embedding, chỉ mục, truy xuất hoặc số đo tác động trong luồng ứng dụng.')
        self.p('Công cụ AI được sử dụng để hỗ trợ rà soát mã nguồn, thu thập tài liệu, xây dựng công cụ đánh giá và biên soạn bản thảo. Sinh viên cần đọc, kiểm chứng và chịu trách nhiệm về bản nộp, đồng thời bổ sung những thay đổi phát sinh sau mốc cập nhật. Bản này không thay cho lời cam đoan và chữ ký xác nhận của sinh viên.')
        self.h('TÓM TẮT',1)
        self.p('Đề tài xây dựng PhishLens, một ứng dụng hỗ trợ kiểm tra nguy cơ lừa đảo trong URL, email và tin nhắn tiếng Việt. Hệ thống kết hợp mô hình phân loại URL dựa trên đặc trưng chuỗi với mô hình ngôn ngữ lớn để phân tích hành động được yêu cầu và giải thích kết quả. Phần mở rộng RAG hướng tới bổ sung kiến thức cảnh báo và quy trình hợp lệ có nguồn đối chiếu.')
        self.p('Bản hiện tại sử dụng XGBoost với 21 đặc trưng URL, DeepSeek-V4-Flash qua API FPT và giao diện web cục bộ. Trên 8.941 URL kiểm tra PhishVN, ngưỡng cảnh báo đạt Precision 81,83%, Recall 99,39%, F1 89,76% và tỷ lệ báo nhầm 29,26%. Trên mẫu cân bằng 100 URL, nhận định LLM khi được cung cấp kết quả URL đạt F1 90,09%, so với 63,16% của LLM độc lập; tuy nhiên tỷ lệ báo nhầm tăng từ 4% lên 22%. Quy tắc kết hợp hiện hành đưa F1 đầu ra cuối về 88,50%, cho thấy thêm thành phần chưa đồng nghĩa với kết quả cuối tốt hơn.')
        self.p('Dự án đã chuẩn bị 25 bản tóm lược tiếng Việt có dẫn nguồn từ 11 đơn vị và có 60 ca kiểm thử tự động đạt. Các kết quả hiện tại chủ yếu phản ánh tác vụ URL; chưa chứng minh hiệu quả trên email thực tế hoặc tác dụng của RAG. Công việc tiếp theo tập trung vào truy xuất quy mô nhỏ, tập thông điệp được gán nhãn độc lập, kiểm tra nguồn dẫn và giảm báo nhầm mà không che giấu các trường hợp thiếu thông tin.')
        self.p('Từ khóa: phishing; PhishVN; XGBoost; mô hình ngôn ngữ lớn; RAG.',small=True)
        self.h('MỤC LỤC',1);p=self.doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0)
        field(p,' TOC \\o "1-2" \\h \\z \\u ','Mục lục tự động')
        self.h('DANH MỤC HÌNH VẼ VÀ BẢNG BIỂU',1)
        self.list_anchor=self.doc.add_paragraph()
        self.h('DANH MỤC TỪ VIẾT TẮT',1)
        self.table('Các thuật ngữ dùng trong báo cáo',['Viết tắt','Tiếng Anh','Diễn giải'],[
            ['API','Application Programming Interface','Giao diện lập trình ứng dụng'],['LLM','Large Language Model','Mô hình ngôn ngữ lớn'],['RAG','Retrieval-Augmented Generation','Sinh nội dung có bổ sung tài liệu truy xuất'],['ML','Machine Learning','Học máy'],['URL','Uniform Resource Locator','Địa chỉ tài nguyên'],['OTP','One-Time Password','Mã xác thực dùng một lần'],['FPR','False Positive Rate','Tỷ lệ báo nhầm trên lớp âm'],['TP / FP','True / False Positive','Cảnh báo đúng / cảnh báo nhầm'],['TN / FN','True / False Negative','Âm đúng / bỏ sót'],['PII','Personally Identifiable Information','Thông tin nhận dạng cá nhân'],['p95','95th percentile','Phân vị 95 của thời gian đo'],['UNKNOWN','Abstention / insufficient evidence','Chưa đủ thông tin để kết luận']],widths=[2.3,5.7,8])
        # This is an unnumbered terminology table, not part of the chapter table list.
        self.captions.pop()
        self.section()
        self.frontmode=False
    def finish(self):
        self.h('TÀI LIỆU THAM KHẢO',1)
        self.p('Các nguồn trực tuyến được đối chiếu ngày 26/10/2026. Số đo của đề tài lấy từ tệp kết quả cục bộ được liệt kê ở Phụ lục C; liên kết GitHub là nơi công bố mã nguồn và không thay thế bằng chứng của bản làm việc chưa đẩy lên remote.',small=True)
        for key,n in self.refs.items():
            author,title,date,url=REFS[key]
            p=self.p(f'[{n}] {author}, “{title}”'+(f', {date}.' if date else '.')+' [Trực tuyến].',small=True)
            p.paragraph_format.first_line_indent=Cm(-.7);p.paragraph_format.left_indent=Cm(.7)
            p=self.doc.add_paragraph();p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_after=Pt(9)
            hyperlink(p,url,url,10)
        # Insert the caption list into the front matter using live page references.
        anchor=self.list_anchor._p
        for kind in ['Hình','Bảng']:
            p=OxmlElement('w:p');anchor.addnext(p);anchor=p
            from docx.text.paragraph import Paragraph
            pp=Paragraph(p,self.list_anchor._parent);pp.add_run('Danh mục '+('hình vẽ' if kind=='Hình' else 'bảng biểu')).bold=True
            pp.paragraph_format.first_line_indent=Cm(0)
            for typ,text,bm in self.captions:
                if typ!=kind:continue
                p=OxmlElement('w:p');anchor.addnext(p);anchor=p;pp=Paragraph(p,self.list_anchor._parent)
                pp.paragraph_format.first_line_indent=Cm(0);pp.paragraph_format.space_after=Pt(2)
                pp.paragraph_format.line_spacing=1.05
                rr=pp.add_run(text+'  ·  ');rr.font.size=Pt(11);field(pp,f' PAGEREF {bm} \\h ')
        self.doc.save(OUTPUT/(NAME+'.docx'))

def diagrams():
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10,'savefig.facecolor':'white'})
    def make(name,size=(12,7)):
        fig,ax=plt.subplots(figsize=size);ax.set_xlim(0,12);ax.set_ylim(0,7);ax.axis('off');return fig,ax
    def box(ax,x,y,w,h,s,dashed=False,color='#edf3f7'):
        ax.add_patch(FancyBboxPatch((x,y),w,h,boxstyle='round,pad=0.09,rounding_size=0.10',edgecolor='#35546b',facecolor=color,linewidth=1.3,linestyle='--' if dashed else '-'))
        ax.text(x+w/2,y+h/2,s,ha='center',va='center',fontsize=10)
    def arrow(ax,a,b,dashed=False):ax.add_patch(FancyArrowPatch(a,b,arrowstyle='-|>',mutation_scale=13,color='#35546b',linewidth=1.2,linestyle='--' if dashed else '-'))
    fig,ax=make('architecture')
    box(ax,.2,5.4,2.5,1,'Người dùng\nTrình duyệt cục bộ')
    box(ax,3.5,5.4,3,1,'FastAPI\nKiểm tra đầu vào')
    arrow(ax,(2.8,5.9),(3.4,5.9))
    box(ax,.2,3.1,2.6,1.2,'URL được trích xuất\n21 đặc trưng → XGBoost')
    box(ax,3.5,3.1,3,1.2,'Che dữ liệu → xem trước\nNgười dùng xác nhận')
    box(ax,8,3.1,3.5,1.2,'FPT API\nDeepSeek-V4-Flash')
    arrow(ax,(4.3,5.3),(1.5,4.4));arrow(ax,(5.1,5.3),(5.1,4.4));arrow(ax,(6.6,3.7),(7.9,3.7))
    ax.plot([2.3,2.3,9.1],[4.3,4.7,4.7],color='#35546b',linewidth=1.1)
    arrow(ax,(9.1,4.7),(9.1,4.4))
    ax.text(6.6,4.7,'Kết quả ML đã che, nếu bật',ha='center',va='center',fontsize=8.5,bbox={'facecolor':'white','edgecolor':'none','pad':2})
    box(ax,8,5.4,3.5,1,'Kho 25 nguồn → embedding\nFAISS → top-k tài liệu',True,'#fff4df')
    arrow(ax,(9.7,5.3),(9.7,4.4),True)
    box(ax,3.5,.7,4.2,1.2,'Kiểm tra JSON và bằng chứng\nFusion → cảnh báo + xuất JSON')
    arrow(ax,(1.6,3),(3.8,1.9));arrow(ax,(9.7,3),(7.1,1.9))
    ax.text(8.2,.9,'Nét liền: đã có\nNét đứt: thiết kế RAG',fontsize=9,color='#725d33')
    fig.savefig(ASSETS/'kien_truc.png',dpi=220,bbox_inches='tight');plt.close(fig)
    fig,ax=make('usecase',(12,5.5))
    box(ax,.1,2.7,2,1,'Người dùng')
    for x,y,s,d in [(3.1,5.3,'Nhập URL / email / tin nhắn',False),(3.1,3.6,'Xem và xác nhận bản gửi API',False),(3.1,1.9,'Xem kết quả và dấu hiệu',False),(7.4,1.9,'Tải kết quả JSON',False),(7.4,5.3,'Truy xuất và đối chiếu nguồn',True),(7.4,3.6,'Cấu hình kết nối LLM',False)]:
        box(ax,x,y,3.7,1,s,d)
        if x==3.1:arrow(ax,(2.2,3.2),(x-.1,y+.5))
    arrow(ax,(6.9,2.4),(7.3,2.4));arrow(ax,(6.9,5.8),(7.3,5.8),True)
    ax.text(.2,.7,'Quản trị kho bằng script; chưa có tài khoản hay giao diện quản trị.',fontsize=10)
    fig.savefig(ASSETS/'use_case.png',dpi=220,bbox_inches='tight');plt.close(fig)
    fig,ax=make('rag')
    steps=[(.2,5.4,'25 bản tóm lược có nguồn\nKiểm tra hash / phiên bản'),(4.2,5.4,'Tokenizer → chia đoạn\npassage: → embedding'),(8.2,5.4,'Chuẩn hóa vector\nFAISS IndexFlatIP'),(.2,2.9,'Nội dung đã che\nquery: → embedding'),(4.2,2.9,'Top 3 → lọc ngưỡng\nTối đa 1 đoạn / nguồn'),(8.2,2.9,'LLM: dữ liệu đầu vào\n+ kiến thức tham khảo')]
    for x,y,s in steps:box(ax,x,y,3.3,1.2,s,True,'#fff4df')
    for a,b in [((3.6,6),(4.1,6)),((7.6,6),(8.1,6)),((3.6,3.5),(4.1,3.5)),((7.6,3.5),(8.1,3.5)),((9.8,5.3),(6.4,4.2))]:arrow(ax,a,b,True)
    box(ax,4.2,.6,7.3,1.2,'Kiểm tra quote từ đầu vào; citation nằm trong tập truy xuất\nKiểm tra nguồn có hỗ trợ phát biểu; không suy ra blacklist',True,'#fff4df')
    arrow(ax,(9.8,2.8),(8.7,1.9),True)
    ax.text(.2,.8,'Toàn bộ luồng này\nchưa tích hợp runtime.',fontsize=10,color='#725d33')
    fig.savefig(ASSETS/'rag_du_kien.png',dpi=220,bbox_inches='tight');plt.close(fig)
    metrics=json.loads((ROOT/'artifacts/evaluations/phishvn_official_v1/url_model_metrics.json').read_text())['test']
    fig,axes=plt.subplots(1,2,figsize=(11,4.5))
    names=['WARNING','HIGH','0,5'];xs=range(3)
    for field_,label,color in [('precision','Precision','#3b6b8d'),('recall','Recall','#2c8b78'),('f1','F1','#bd8d39')]:
        yy=[metrics[k][field_]*100 for k in ['warning','high','fixed_0.5']];axes[0].plot(xs,yy,'o-',label=label,color=color)
    axes[0].set_xticks(list(xs),names);axes[0].set_ylim(0,105);axes[0].set_ylabel('%');axes[0].set_title('Đánh đổi khi dùng các ngưỡng đã khóa');axes[0].legend(loc='lower left');axes[0].grid(alpha=.2)
    vals=[metrics[k]['false_positive_rate']*100 for k in ['warning','high','fixed_0.5']]
    axes[1].bar(names,vals,color=['#bf6a59','#bf9c57','#608c89'])
    for i,v in enumerate(vals):axes[1].text(i,v+.6,f'{v:.2f}%',ha='center')
    axes[1].set_ylim(0,36);axes[1].set_ylabel('FPR (%)');axes[1].set_title('Báo nhầm trên 3.845 URL benign');axes[1].grid(axis='y',alpha=.2)
    fig.tight_layout();fig.savefig(ASSETS/'danh_doi_nguong.png',dpi=220);plt.close(fig)

if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--out',type=Path,default=OUTPUT,help='Build destination; published report snapshots are never overwritten.')
    args=parser.parse_args()
    OUTPUT=args.out.resolve()
    if OUTPUT.is_relative_to(ROOT/'docs/reports'):
        parser.error('Choose an output directory outside the published report snapshots.')
    ASSETS=OUTPUT/'assets'
    ASSETS.mkdir(parents=True,exist_ok=True)
    (OUTPUT/'evidence').mkdir(exist_ok=True)
    for name in ['giao_dien.png','xac_nhan_du_lieu.png','ket_qua_ml.png']:
        shutil.copy2(REPORT_SOURCE/'assets'/name,ASSETS/name)
    from report_content import compose, appendices
    diagrams();b=Report();b.front();compose(b);b.finish();appendices(b)
    b.doc.save(OUTPUT/(NAME+'.docx'))
    manifest={str(p.relative_to(ROOT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in [ROOT/'app/fusion.py',ROOT/'app/llm_service.py',ROOT/'models/inference_XGBoost_d6_s42/config.json',ROOT/'artifacts/evaluations/phishvn_official_v1/summary.json',ROOT/'artifacts/evaluations/phishvn_official_v1/url_model_metrics.json',ROOT/'data/knowledge/corpus_v1/documents.jsonl']}
    (OUTPUT/'evidence/report_inputs.json').write_text(json.dumps(manifest,indent=2)+'\n')
    print('Created',OUTPUT/(NAME+'.docx'))
    print('References',len(b.refs),'captions',len(b.captions))
